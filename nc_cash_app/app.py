"""
Walk-On's Cash Tracker (multi-account)
Upload a .qbo file per account, it auto-detects the account by its bank
account number, parses + dedupes transactions into SQLite, and lets you
download an always-current Excel export per account.
"""
import hmac
import io
import os
import re
import sqlite3
from datetime import datetime

from flask import (
    Flask, request, render_template, send_file, redirect, url_for, flash,
    session, Response
)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

APP_DIR = os.path.dirname(os.path.abspath(__file__))
DB_DIR = os.environ.get("DB_DIR", APP_DIR)
DB_PATH = os.path.join(DB_DIR, "cash_tracker.db")

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-change-me")

CATEGORIES = ["Sales", "COGS", "Payroll", "Vendor/AP", "Construction/CapEx",
              "Transfer", "Financing", "Fees", "Other"]

# When an uploaded bank transaction matches a still-open Projected transaction
# on amount (exact) and falls within this many days of it, it gets flagged as
# a possible match for the user to confirm rather than silently merged.
MATCH_WINDOW_DAYS = 21

SCHEMA_VERSION = "3"

# ---------------------------------------------------------------------------
# Password protection (HTTP Basic Auth over the whole app)
# ---------------------------------------------------------------------------
# Credentials come from environment variables set on Render -- never hardcode
# them here or commit them to GitHub. If APP_PASSWORD isn't set, the app
# fails CLOSED (blocks everyone) rather than silently staying open.
APP_USERNAME = os.environ.get("APP_USERNAME", "")
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")


def _check_credentials(username, password):
    if not APP_USERNAME or not APP_PASSWORD:
        return False
    user_ok = hmac.compare_digest(username or "", APP_USERNAME)
    pass_ok = hmac.compare_digest(password or "", APP_PASSWORD)
    return user_ok and pass_ok


def _auth_challenge():
    return Response(
        "Authentication required.", 401,
        {"WWW-Authenticate": "Basic realm=\"Walk-On's Cash Tracker\""},
    )


@app.before_request
def require_login():
    auth = request.authorization
    if not auth or not _check_credentials(auth.username, auth.password):
        return _auth_challenge()


# ---------------------------------------------------------------------------
# Database
# ---------------------------------------------------------------------------

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def _table_exists(conn, name):
    return conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name=?", (name,)
    ).fetchone() is not None


def _column_exists(conn, table, column):
    cols = [r["name"] for r in conn.execute(f"PRAGMA table_info({table})")]
    return column in cols


def init_db():
    conn = get_db()

    conn.executescript("""
    CREATE TABLE IF NOT EXISTS meta (
        key TEXT PRIMARY KEY,
        value TEXT
    );
    """)
    version_row = conn.execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()
    current_version = version_row["value"] if version_row else None

    if current_version == SCHEMA_VERSION:
        conn.close()
        return

    if not _table_exists(conn, "accounts"):
        conn.executescript("""
        CREATE TABLE accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            starting_balance REAL NOT NULL DEFAULT 0,
            as_of_date TEXT NOT NULL,
            last4 TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );
        """)

    legacy_settings = None
    if _table_exists(conn, "settings"):
        rows = conn.execute("SELECT key, value FROM settings").fetchall()
        legacy_settings = {r["key"]: r["value"] for r in rows}

    if legacy_settings and conn.execute("SELECT COUNT(*) c FROM accounts").fetchone()["c"] == 0:
        cur = conn.execute(
            "INSERT INTO accounts (name, starting_balance, as_of_date, last4) "
            "VALUES (?, ?, ?, ?)",
            (
                legacy_settings.get("account_name", "Account 1"),
                float(legacy_settings.get("starting_balance", "0")),
                legacy_settings.get("as_of_date", datetime.now().date().isoformat()),
                legacy_settings.get("account_last4", ""),
            ),
        )
        legacy_account_id = cur.lastrowid
    else:
        legacy_account_id = None

    if not _table_exists(conn, "transactions"):
        conn.executescript("""
        CREATE TABLE transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            account_id INTEGER NOT NULL REFERENCES accounts(id),
            txn_date TEXT NOT NULL,
            txn_type TEXT NOT NULL,
            payee TEXT,
            description TEXT,
            category TEXT,
            amount REAL NOT NULL,
            bank_id TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(account_id, bank_id)
        );
        """)
    elif not _column_exists(conn, "transactions", "account_id"):
        old_rows = conn.execute(
            "SELECT txn_date, txn_type, payee, description, category, amount, "
            "bank_id, created_at FROM transactions"
        ).fetchall()
        conn.execute("ALTER TABLE transactions RENAME TO transactions_old")
        conn.executescript("""
        CREATE TABLE transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            account_id INTEGER NOT NULL REFERENCES accounts(id),
            txn_date TEXT NOT NULL,
            txn_type TEXT NOT NULL,
            payee TEXT,
            description TEXT,
            category TEXT,
            amount REAL NOT NULL,
            bank_id TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(account_id, bank_id)
        );
        """)
        target_account_id = legacy_account_id or conn.execute(
            "SELECT id FROM accounts ORDER BY id LIMIT 1"
        ).fetchone()["id"]
        for r in old_rows:
            conn.execute(
                "INSERT INTO transactions (account_id, txn_date, txn_type, payee, "
                "description, category, amount, bank_id, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (target_account_id, r["txn_date"], r["txn_type"], r["payee"],
                 r["description"], r["category"], r["amount"], r["bank_id"],
                 r["created_at"]),
            )
        conn.execute("DROP TABLE transactions_old")

    if not _column_exists(conn, "transactions", "matched_projected_id"):
        conn.execute(
            "ALTER TABLE transactions ADD COLUMN matched_projected_id "
            "INTEGER REFERENCES transactions(id)"
        )

    if _table_exists(conn, "settings"):
        conn.execute("DROP TABLE settings")

    if conn.execute("SELECT COUNT(*) c FROM accounts").fetchone()["c"] == 0:
        conn.execute(
            "INSERT INTO accounts (name, starting_balance, as_of_date, last4) "
            "VALUES (?, ?, ?, ?)",
            ("NC KKBM LLC", 50.00, "2026-09-07", "4142"),
        )

    conn.execute(
        "INSERT INTO meta (key, value) VALUES ('schema_version', ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (SCHEMA_VERSION,),
    )
    conn.commit()
    conn.close()


def get_accounts():
    conn = get_db()
    rows = conn.execute("SELECT * FROM accounts ORDER BY name ASC").fetchall()
    conn.close()
    return rows


def get_account(account_id):
    conn = get_db()
    row = conn.execute("SELECT * FROM accounts WHERE id=?", (account_id,)).fetchone()
    conn.close()
    return row


def get_all_account_balances():
    conn = get_db()
    accounts = conn.execute("SELECT * FROM accounts ORDER BY name COLLATE NOCASE ASC").fetchall()
    results = []
    for a in accounts:
        total = conn.execute(
            "SELECT COALESCE(SUM(amount), 0) s FROM transactions WHERE account_id=?",
            (a["id"],),
        ).fetchone()["s"]
        results.append({
            "id": a["id"],
            "name": a["name"],
            "last4": a["last4"],
            "as_of_date": a["as_of_date"],
            "balance": a["starting_balance"] + total,
        })
    conn.close()
    return results


def account_balance(account_id):
    conn = get_db()
    acct = conn.execute("SELECT starting_balance FROM accounts WHERE id=?", (account_id,)).fetchone()
    if not acct:
        conn.close()
        return 0.0
    total = conn.execute(
        "SELECT COALESCE(SUM(amount), 0) s FROM transactions WHERE account_id=?",
        (account_id,),
    ).fetchone()["s"]
    conn.close()
    return acct["starting_balance"] + total


# ---------------------------------------------------------------------------
# QBO / OFX parsing
# ---------------------------------------------------------------------------

def parse_ofx_date(raw):
    raw = raw.strip()[:8]
    return datetime.strptime(raw, "%Y%m%d").date().isoformat()


def parse_qbo(file_bytes):
    text = file_bytes.decode("utf-8", errors="replace")

    def find_one(tag, blob):
        m = re.search(rf"<{tag}>([^\r\n<]*)", blob)
        return m.group(1).strip() if m else None

    acctid = find_one("ACCTID", text)
    bankid = find_one("BANKID", text)

    transactions = []
    for block in re.findall(r"<STMTTRN>(.*?)</STMTTRN>", text, re.S):
        dtposted = find_one("DTPOSTED", block)
        trnamt = find_one("TRNAMT", block)
        fitid = find_one("FITID", block)
        name = find_one("NAME", block) or find_one("N", block)
        memo = find_one("MEMO", block)
        if not (dtposted and trnamt and fitid):
            continue
        transactions.append({
            "date": parse_ofx_date(dtposted),
            "amount": float(trnamt),
            "fitid": fitid.strip(),
            "name": (name or "").strip(),
            "memo": (memo or "").strip(),
        })

    return {"acctid": acctid, "bankid": bankid, "transactions": transactions}


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def build_excel(account_id):
    FONT = "Arial"
    HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
    HEADER_FONT = Font(name=FONT, size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F4E78")
    LABEL_FONT = Font(name=FONT, size=10, bold=True)
    INPUT_FONT = Font(name=FONT, size=11, bold=True, color="0000FF")
    NOTE_FONT = Font(name=FONT, size=9, italic=True, color="808080")
    THIN = Side(style="thin", color="B7B7B7")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ACTUAL_FILL = PatternFill("solid", fgColor="E2EFDA")
    PROJECTED_FILL = PatternFill("solid", fgColor="FFF2CC")

    acct = get_account(account_id)
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM transactions WHERE account_id=? ORDER BY txn_date ASC, id ASC",
        (account_id,),
    ).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    setup = wb.active
    setup.title = "Setup"
    setup.sheet_view.showGridLines = False

    setup["B2"] = f"{acct['name']} — Cash Tracker Setup"
    setup["B2"].font = TITLE_FONT
    setup.merge_cells("B2:D2")

    setup["B4"] = "Account Name"; setup["B4"].font = LABEL_FONT
    setup["C4"] = acct["name"]; setup["C4"].font = INPUT_FONT

    setup["B5"] = "Starting Balance"; setup["B5"].font = LABEL_FONT
    setup["C5"] = acct["starting_balance"]; setup["C5"].font = INPUT_FONT
    setup["C5"].number_format = "$#,##0.00"

    setup["B6"] = "As-Of Date"; setup["B6"].font = LABEL_FONT
    setup["C6"] = acct["as_of_date"]; setup["C6"].font = INPUT_FONT

    setup["B7"] = "Account Number (last 4)"; setup["B7"].font = LABEL_FONT
    setup["C7"] = acct["last4"]; setup["C7"].font = INPUT_FONT

    setup["B8"] = "Current Cash Balance"; setup["B8"].font = LABEL_FONT
    running = acct["starting_balance"]
    for r in rows:
        running += r["amount"]
    setup["C8"] = running
    setup["C8"].number_format = "$#,##0.00"
    setup["C8"].font = Font(name=FONT, size=12, bold=True)

    setup["B10"] = "Generated by the Walk-On's Cash Tracker app — this file is a live export, re-download after any change."
    setup["B10"].font = NOTE_FONT

    for col, width in [("A", 3), ("B", 30), ("C", 20), ("D", 20)]:
        setup.column_dimensions[col].width = width

    ledger = wb.create_sheet("Ledger")
    ledger.sheet_view.showGridLines = False
    headers = ["Date", "Type", "Payee/Payor", "Description", "Category",
               "Amount", "Running Balance", "Bank ID (dedup key)"]
    for i, h in enumerate(headers, start=1):
        c = ledger.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ledger.freeze_panes = "A2"
    widths = [12, 12, 22, 30, 16, 14, 16, 26]
    for i, w in enumerate(widths, start=1):
        ledger.column_dimensions[get_column_letter(i)].width = w

    running = acct["starting_balance"]
    r_idx = 2
    for r in rows:
        running += r["amount"]
        vals = [
            datetime.fromisoformat(r["txn_date"]).date(),
            r["txn_type"], r["payee"], r["description"], r["category"],
            r["amount"], running, r["bank_id"] or "",
        ]
        fill = ACTUAL_FILL if r["txn_type"] == "Actual" else PROJECTED_FILL
        for c_idx, v in enumerate(vals, start=1):
            cell = ledger.cell(row=r_idx, column=c_idx, value=v)
            cell.border = BORDER
            cell.fill = fill
            cell.font = Font(name=FONT, size=10)
            if c_idx == 1:
                cell.number_format = "mm/dd/yyyy"
            if c_idx in (6, 7):
                cell.number_format = "$#,##0.00;($#,##0.00)"
        r_idx += 1

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, acct["name"]


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.route("/", methods=["GET"])
def dashboard():
    accounts = get_all_account_balances()
    if not accounts:
        return redirect(url_for("manage_accounts"))
    return render_template("dashboard.html", accounts=accounts)


@app.route("/account/<int:account_id>", methods=["GET"])
def account_page(account_id):
    acct = get_account(account_id)
    if not acct:
        flash("Account not found.", "error")
        return redirect(url_for("dashboard"))

    session["account_id"] = account_id
    accounts = get_accounts()

    conn = get_db()
    all_rows = conn.execute(
        "SELECT id, amount FROM transactions WHERE account_id=? ORDER BY txn_date ASC, id ASC",
        (account_id,),
    ).fetchall()
    running = acct["starting_balance"]
    running_by_id = {}
    for r in all_rows:
        running += r["amount"]
        running_by_id[r["id"]] = running

    display_rows = conn.execute(
        "SELECT * FROM transactions WHERE account_id=? ORDER BY txn_date DESC, id DESC LIMIT 100",
        (account_id,),
    ).fetchall()

    match_ids = [r["matched_projected_id"] for r in display_rows if r["matched_projected_id"]]
    match_lookup = {}
    if match_ids:
        placeholders = ",".join("?" for _ in match_ids)
        for m in conn.execute(
            f"SELECT id, txn_date, payee, description, amount FROM transactions "
            f"WHERE id IN ({placeholders})",
            match_ids,
        ).fetchall():
            match_lookup[m["id"]] = m
    conn.close()

    rows = []
    for r in display_rows:
        row = dict(r, running_total=running_by_id[r["id"]])
        row["match_candidate"] = (
            match_lookup.get(r["matched_projected_id"]) if r["matched_projected_id"] else None
        )
        rows.append(row)

    return render_template(
        "index.html",
        rows=rows,
        balance=account_balance(account_id),
        accounts=accounts,
        current_account=acct,
        categories=CATEGORIES,
    )


@app.route("/upload", methods=["POST"])
def upload():
    f = request.files.get("qbo_file")
    if not f or f.filename == "":
        flash("No file selected.", "error")
        return redirect(url_for("dashboard"))

    data = parse_qbo(f.read())
    acctid = data.get("acctid") or ""
    found_last4 = acctid[-4:] if len(acctid) >= 4 else acctid

    accounts = get_accounts()
    matches = [a for a in accounts if a["last4"] and a["last4"] == found_last4]

    if not matches:
        flash(
            f"This file is for account ending {found_last4 or '(unknown)'}, which "
            f"isn't set up yet. Add it under Manage Accounts first, then re-upload.",
            "error",
        )
        return redirect(url_for("dashboard"))

    acct = matches[0]
    as_of_date = acct["as_of_date"]

    conn = get_db()
    inserted, dup_skipped, historical_skipped, matched = 0, 0, 0, 0
    for t in data["transactions"]:
        if t["date"] <= as_of_date:
            historical_skipped += 1
            continue
        try:
            cur = conn.execute(
                "INSERT INTO transactions (account_id, txn_date, txn_type, payee, "
                "description, category, amount, bank_id) "
                "VALUES (?, ?, 'Actual', ?, ?, ?, ?, ?)",
                (acct["id"], t["date"], t["name"], t["memo"], "", t["amount"], t["fitid"]),
            )
            inserted += 1
            new_id = cur.lastrowid

            # Flag a possible match against a still-open Projected transaction:
            # same account, exact amount, within MATCH_WINDOW_DAYS of this
            # transaction's date, and not already claimed by another match.
            candidate = conn.execute(
                "SELECT id FROM transactions "
                "WHERE account_id=? AND txn_type='Projected' AND amount=? "
                "AND id NOT IN (SELECT matched_projected_id FROM transactions "
                "WHERE matched_projected_id IS NOT NULL) "
                "AND ABS(JULIANDAY(txn_date) - JULIANDAY(?)) <= ? "
                "ORDER BY ABS(JULIANDAY(txn_date) - JULIANDAY(?)) ASC LIMIT 1",
                (acct["id"], t["amount"], t["date"], MATCH_WINDOW_DAYS, t["date"]),
            ).fetchone()
            if candidate:
                conn.execute(
                    "UPDATE transactions SET matched_projected_id=? WHERE id=?",
                    (candidate["id"], new_id),
                )
                matched += 1
        except sqlite3.IntegrityError:
            dup_skipped += 1
    conn.commit()
    conn.close()

    session["account_id"] = acct["id"]
    msg = (f"Matched account '{acct['name']}' (ending {found_last4}). "
           f"Imported {inserted} new transaction(s), skipped {dup_skipped} duplicate(s)")
    if matched:
        msg += f", flagged {matched} possible match{'es' if matched != 1 else ''} with projected transactions for you to confirm"
    if historical_skipped:
        msg += f", skipped {historical_skipped} dated on/before the {as_of_date} starting balance"
    msg += "."
    flash(msg, "success" if inserted else "info")
    return redirect(url_for("account_page", account_id=acct["id"]))


@app.route("/add", methods=["POST"])
def add_transaction():
    account_id = request.form.get("account_id", type=int)
    txn_date = request.form.get("txn_date")
    payee = request.form.get("payee", "").strip()
    description = request.form.get("description", "").strip()
    category = request.form.get("category", "").strip()
    amount = request.form.get("amount", "").strip()

    if not account_id or not get_account(account_id):
        flash("Select a valid account first.", "error")
        return redirect(url_for("dashboard"))

    try:
        amount_val = float(amount)
        datetime.strptime(txn_date, "%Y-%m-%d")
    except (ValueError, TypeError):
        flash("Please provide a valid date and numeric amount.", "error")
        return redirect(url_for("account_page", account_id=account_id))

    conn = get_db()
    conn.execute(
        "INSERT INTO transactions (account_id, txn_date, txn_type, payee, description, "
        "category, amount, bank_id) VALUES (?, ?, 'Projected', ?, ?, ?, ?, NULL)",
        (account_id, txn_date, payee, description, category, amount_val),
    )
    conn.commit()
    conn.close()
    flash("Projected transaction added.", "success")
    return redirect(url_for("account_page", account_id=account_id))


@app.route("/delete/<int:txn_id>", methods=["POST"])
def delete_transaction(txn_id):
    conn = get_db()
    row = conn.execute(
        "SELECT txn_type, account_id FROM transactions WHERE id=?", (txn_id,)
    ).fetchone()
    account_id = row["account_id"] if row else None
    if row and row["txn_type"] == "Projected":
        conn.execute("DELETE FROM transactions WHERE id=?", (txn_id,))
        conn.commit()
        flash("Projected transaction removed.", "success")
    else:
        flash("Only Projected rows can be deleted here.", "error")
    conn.close()
    return redirect(url_for("account_page", account_id=account_id))


@app.route("/match/<int:txn_id>/confirm", methods=["POST"])
def confirm_match(txn_id):
    conn = get_db()
    row = conn.execute(
        "SELECT account_id, matched_projected_id FROM transactions WHERE id=?",
        (txn_id,),
    ).fetchone()
    if not row or not row["matched_projected_id"]:
        flash("No pending match found for that transaction.", "error")
        conn.close()
        return redirect(url_for("dashboard"))

    account_id = row["account_id"]
    # Clear the reference before deleting the projected row it points to,
    # since matched_projected_id is a foreign key back into this same table.
    conn.execute("UPDATE transactions SET matched_projected_id=NULL WHERE id=?", (txn_id,))
    conn.execute("DELETE FROM transactions WHERE id=?", (row["matched_projected_id"],))
    conn.commit()
    conn.close()
    flash("Match confirmed — the projected transaction was removed.", "success")
    return redirect(url_for("account_page", account_id=account_id))


@app.route("/match/<int:txn_id>/dismiss", methods=["POST"])
def dismiss_match(txn_id):
    conn = get_db()
    row = conn.execute(
        "SELECT account_id FROM transactions WHERE id=?", (txn_id,)
    ).fetchone()
    if not row:
        flash("Transaction not found.", "error")
        conn.close()
        return redirect(url_for("dashboard"))

    account_id = row["account_id"]
    conn.execute("UPDATE transactions SET matched_projected_id=NULL WHERE id=?", (txn_id,))
    conn.commit()
    conn.close()
    flash("Dismissed — no changes made.", "info")
    return redirect(url_for("account_page", account_id=account_id))


@app.route("/download/<int:account_id>")
def download(account_id):
    if not get_account(account_id):
        flash("Account not found.", "error")
        return redirect(url_for("dashboard"))
    buf, account_name = build_excel(account_id)
    filename = f"{account_name.replace(' ', '_')}_Cash_Tracker.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/accounts", methods=["GET"])
def manage_accounts():
    return render_template("accounts.html", accounts=get_accounts())


@app.route("/accounts/new", methods=["POST"])
def create_account():
    name = request.form.get("name", "").strip()
    starting_balance = request.form.get("starting_balance", "0").strip()
    as_of_date = request.form.get("as_of_date", "").strip()
    last4 = request.form.get("last4", "").strip()

    try:
        starting_balance_val = float(starting_balance)
        datetime.strptime(as_of_date, "%Y-%m-%d")
    except (ValueError, TypeError):
        flash("Please provide a valid starting balance and as-of date.", "error")
        return redirect(url_for("manage_accounts"))

    if not name:
        flash("Account name is required.", "error")
        return redirect(url_for("manage_accounts"))

    conn = get_db()
    cur = conn.execute(
        "INSERT INTO accounts (name, starting_balance, as_of_date, last4) VALUES (?, ?, ?, ?)",
        (name, starting_balance_val, as_of_date, last4),
    )
    new_id = cur.lastrowid
    conn.commit()
    conn.close()
    session["account_id"] = new_id
    flash(f"Account '{name}' added.", "success")
    return redirect(url_for("account_page", account_id=new_id))


@app.route("/accounts/<int:account_id>/update", methods=["POST"])
def update_account(account_id):
    if not get_account(account_id):
        flash("Account not found.", "error")
        return redirect(url_for("manage_accounts"))

    name = request.form.get("name", "").strip()
    starting_balance = request.form.get("starting_balance", "").strip()
    as_of_date = request.form.get("as_of_date", "").strip()
    last4 = request.form.get("last4", "").strip()

    try:
        starting_balance_val = float(starting_balance)
        datetime.strptime(as_of_date, "%Y-%m-%d")
    except (ValueError, TypeError):
        flash("Please provide a valid starting balance and as-of date.", "error")
        return redirect(url_for("manage_accounts"))

    conn = get_db()
    conn.execute(
        "UPDATE accounts SET name=?, starting_balance=?, as_of_date=?, last4=? WHERE id=?",
        (name, starting_balance_val, as_of_date, last4, account_id),
    )
    conn.commit()
    conn.close()
    flash(f"'{name}' updated.", "success")
    return redirect(url_for("manage_accounts"))


init_db()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
